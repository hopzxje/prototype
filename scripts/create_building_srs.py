from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


REFERENCE = Path(r"C:\Users\tienhop\Downloads\Report-3_Software Requirement Specification (4).docx")
OUTPUT = Path(r"D:\ki9\prototype\artifacts\Create_Building_Functional_Requirements.docx")
SCREENSHOT = Path(r"D:\ki9\prototype\test.png")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_table_borders(table, color='000000', size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_table_layout(table, widths_twips=(2300, 6900)):
    """Keep text inside a predictable two-column Word table."""
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_w.set(qn('w:w'), str(sum(widths_twips)))

    layout = tbl_pr.first_child_found_in('w:tblLayout')
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tbl_pr.append(layout)
    layout.set(qn('w:type'), 'fixed')

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths_twips):
        grid_col.set(qn('w:w'), str(width))

    cell_mar = tbl_pr.first_child_found_in('w:tblCellMar')
    if cell_mar is None:
        cell_mar = OxmlElement('w:tblCellMar')
        tbl_pr.append(cell_mar)
    for side, width in {'top': 80, 'left': 110, 'bottom': 80, 'right': 110}.items():
        node = cell_mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side)
            cell_mar.append(node)
        node.set(qn('w:w'), str(width))
        node.set(qn('w:type'), 'dxa')

    for row in table.rows:
        for cell, width in zip(row.cells, widths_twips):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:type'), 'dxa')
            tc_w.set(qn('w:w'), str(width))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    set_table_borders(table)
    set_table_layout(table)
    header = table.rows[0].cells
    header[0].width = Inches(1.6)
    header[1].width = Inches(4.85)
    header[0].text = rows[0][0]
    header[1].text = rows[0][1]
    for cell in header:
        set_cell_shading(cell, 'FCE4CC')
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = Pt(9)
    for left, right in rows[1:]:
        cells = table.add_row().cells
        cells[0].width = Inches(1.6)
        cells[1].width = Inches(4.85)
        cells[0].text = left
        cells[1].text = right
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        set_table_layout(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.add_run('- ' + text)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text, style='Normal')
    p.paragraph_format.space_after = Pt(6)
    return p


doc = Document(str(REFERENCE))
body = doc._body._element
sect_pr = body.sectPr
for child in list(body):
    if child is not sect_pr:
        body.remove(child)

# The supplied template stores page margins as decimal twips, which is not
# accepted by python-docx when calculating table width. Normalize the copy's
# page margins to valid Word twips while preserving the template's portrait
# page layout and style definitions.
pg_mar = sect_pr.find(qn('w:pgMar'))
if pg_mar is not None:
    for key, value in {'top': '720', 'bottom': '720', 'left': '720', 'right': '720', 'header': '360', 'footer': '360'}.items():
        pg_mar.set(qn('w:' + key), value)

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal'].font.size = Pt(10)
for name in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
    styles[name].font.name = 'Arial'
    styles[name].font.color.rgb = None

title = doc.add_paragraph('III. Functional Requirements', style='Heading 1')
title.paragraph_format.space_after = Pt(10)
add_body(doc, 'This section specifies the screens and system behavior for creating a new building and preparing its rooms for operation. The flow is available to Administrator and Manager roles. Staff users can view building information but cannot create a building.')

doc.add_paragraph('1. Building Management', style='Heading 2')
doc.add_paragraph('1.1 Create Building', style='Heading 3')

doc.add_paragraph('1.1.1 Building Management Screen', style='Heading 4')
if SCREENSHOT.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SCREENSHOT), width=Inches(6.45))
    cap = doc.add_paragraph('Figure 1. Building Management Screen', style='Normal')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
add_body(doc, 'This screen is the entry point for the Create Building flow. It shows the currently selected building, operating status, manager, address, room statistics, room list, and the Add Building action in the top bar.')
add_body(doc, 'This screen allows the authorized user to:')
for item in [
    'View the current building summary and room occupancy metrics.',
    'Choose a building context from the top-bar building selector.',
    'Open the Create Building form by clicking Add Building.',
    'Open the Add Room form for the selected building after a building has been created.',
]:
    add_bullet(doc, item)
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Building selector', 'Initial values: all active buildings. The selected building becomes the context for room and operational data.'),
    ('Add Building', 'Opens the Create Building form. Visible to Administrator and Manager only.'),
    ('Building summary', 'Displays building name, code, address, responsible manager, status, total rooms, occupied rooms, vacant rooms, and occupancy rate.'),
    ('Room list', 'Displays generated rooms grouped by floor and supports search, floor filtering, and status filtering.'),
    ('Global navigation sidebar', 'Identifies the StayHub product and provides links to operational areas. The current item, Building and Room Management, is highlighted.'),
    ('Top bar', 'Shows the page title and breadcrumb, the current building selector, the Add Building action, and the notification entry point.'),
    ('Building overview panel', 'Shows the building icon, building name, building code, operating status, pause action, address, responsible manager, and Add Room action.'),
    ('Operational metrics', 'Shows total rooms, occupied rooms, vacant rooms, and occupancy rate for the selected building.'),
    ('Search and filter bar', 'Supports keyword search by room or tenant, filtering by floor, filtering by room status, and a displayed-room count.'),
    ('Floor groups', 'Organize rooms by floor and show the number of rooms and the occupied-room ratio for each floor.'),
    ('Room cards', 'Summarize room number, floor, room type, area, rental amount, contract status, contract timeline, tenant or handover status, and the room action area.'),
    ('Room number and floor', 'Primary room identifier, for example P101, with the floor badge shown beside it.'),
    ('Room status', 'Status badge such as Occupied, Vacant, Coming soon, or Maintenance.'),
    ('Room type and area', 'Displays the room classification and area in square meters.'),
    ('Rental amount', 'Shows the monthly rental price and whether it is based on an active contract or an agreed price.'),
    ('Contract timeline', 'Shows contract start date, end date, remaining days or percentage, and the relevant contract code where applicable.'),
    ('Tenant or handover status', 'Shows the current tenant name and phone number, or a handover state when the room is reserved but not yet handed over.'),
])

doc.add_paragraph('1.1.2 Create Building Form', style='Heading 4')
add_body(doc, 'This form creates one operational building and generates its initial room inventory from the floor and room-count parameters. The form is opened from the Add Building action.')
add_body(doc, 'The user must complete the required fields and click Save Building Information. The system validates duplicate building codes and names before saving.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Modal header', 'Displays the title Add New Building, the purpose text for initializing an operational building, and the building icon.'),
    ('Close action', 'Closes the modal without saving any changes.'),
    ('Validation error box', 'Hidden by default. Appears at the top of the form when a required value is missing, a value is out of range, or a duplicate building code or name is detected.'),
    ('Building Code', 'Required. Uppercase identifier, for example SH-CAUGIAY. Must be unique across all buildings.'),
    ('Building Name', 'Required. Human-readable name, for example StayHub Cau Giay. Must be unique across all buildings.'),
    ('Building Address', 'Required. Full operating address of the building.'),
    ('Responsible Manager', 'Optional. Name of the manager responsible for the building. A default manager may be provided.'),
    ('Contact Hotline', 'Optional. Phone number used for building operations and contact.'),
    ('Building structure section', 'Groups the floor and room configuration fields used to generate the initial room inventory.'),
    ('Number of Floors', 'Required integer from 1 to 50. Determines the number of floor groups generated.'),
    ('Rooms per Floor', 'Required integer from 1 to 30. Determines the number of rooms generated on each floor.'),
    ('Room Area', 'Required number from 10 to 300 square meters. Applied as the default area for generated rooms.'),
    ('Default Room Price', 'Required. Default monthly rental price applied to generated rooms.'),
    ('Default Room Type', 'Optional. Initial type applied to every generated room, such as Studio or One Bedroom.'),
    ('Cancel', 'Closes the form without creating a building.'),
    ('Save Building Information', 'Validates the form, creates the building, generates rooms from the selected configuration, and refreshes the building overview.'),
])

doc.add_paragraph('1.1.3 Create Building Validation and Processing', style='Heading 4')
add_body(doc, 'When the user submits the form, the system performs validation before changing data.')
for item in [
    'Reject the request when any required field is empty or outside its allowed range.',
    'Reject the request when the building code already exists, ignoring letter case.',
    'Reject the request when the building name already exists, ignoring letter case.',
    'Create the building in Active status when validation succeeds.',
    'Generate room records using the selected floor count and rooms-per-floor value. Room numbers follow the floor convention, for example P101, P102, P201, and P202.',
    'Save the new building and generated rooms together, then refresh the building selector and current overview.',
    'Show a success notification after the operation completes and keep the user on the Building Management screen.',
]:
    add_bullet(doc, item)
add_body(doc, 'Scenario')
add_table(doc, [
    ('Scenario', 'System behavior'),
    ('Duplicate building code', 'The form remains open and displays a validation error. No building or room is created.'),
    ('Duplicate building name', 'The form remains open and displays a validation error. No building or room is created.'),
    ('Invalid floor or room count', 'The relevant field is rejected and the user must correct it before submission.'),
    ('Successful creation', 'The building is saved as Active, rooms are generated, and the overview is refreshed.'),
])

doc.add_paragraph('1.2 Room Management After Building Creation', style='Heading 3')
doc.add_paragraph('1.2.1 Room Detail Screen', style='Heading 4')
add_body(doc, 'This screen provides the complete operational view of one room, opened by selecting a room card from the building overview. It combines room information, resident and contract information, room images, assigned assets, and room-level actions in one workspace.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Room header', 'Shows the room number, room status badge, building name, building address, and the current room-status selector.'),
    ('Close action', 'Closes the room detail view and returns to the building room list.'),
    ('Room status selector', 'Displays the current status, for example Occupied or Vacant, and allows an authorized user to update it according to the room workflow.'),
    ('Room photo gallery', 'Displays the primary room image and supporting images for the room. The user can update or add room photos.'),
    ('Usable area', 'Displays the room area in square meters.'),
    ('Room type and floor', 'Displays the room classification and the floor where the room is located.'),
    ('Expected monthly rent', 'Displays the expected monthly rental amount for the room.'),
    ('Default deposit', 'Displays the configured default deposit amount for the room.'),
    ('Resident and current occupant card', 'Displays the current resident name, phone number, identity document reference, contract status, contract code, contract period, and dependent information when available.'),
    ('Equipment and assets section', 'Displays the number of assets assigned to the room and provides actions to add an asset or assign the same type of asset to multiple rooms.'),
    ('Asset table', 'Lists asset code, asset name and brand/model, asset type, condition, value, and available actions.'),
    ('Asset condition action', 'Allows an authorized user to update the condition of an assigned asset, such as Good, New, Damaged, or Under Repair.'),
    ('Asset recovery action', 'Removes the asset from the room and returns it to available inventory according to the asset workflow.'),
    ('Contract history', 'Shows the active contract or relevant contract history for the room, including contract code, resident, contract period, status, and amount.'),
    ('Room lock or delete action', 'Provides a restricted room-level action. The system must prevent deletion when the room has an active contract or assigned operational data.'),
])
add_body(doc, 'Room Detail Processing Rules')
for item in [
    'Only the selected room context is displayed; actions must apply to this room only.',
    'A room with an active contract cannot be deleted or made unavailable without completing the required contract workflow.',
    'An asset can be recovered only after the system records the recovery action and updates the room asset count.',
    'Changes to room status, photos, residents, contracts, or assets must refresh the detail view and the building overview.'
]:
    add_bullet(doc, item)

doc.add_paragraph('1.2.2 Add Room Screen', style='Heading 4')
add_body(doc, 'This screen is used when the user needs to add a room after the building has already been created. The building is inherited from the selected top-bar context and is not selected again.')
add_body(doc, 'The system automatically generates the room number from the selected floor and prevents duplicate room numbers within the same building.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Modal header', 'Displays the title Add New Room to Building, the purpose text about following the building floor and room limits, and the add-room icon.'),
    ('Close action', 'Closes the modal without saving the new room.'),
    ('Building', 'Read-only display of the currently selected building context. The user does not select the building again in this form.'),
    ('Floor', 'Required. Selects an existing floor and displays its occupancy, for example 0/4 rooms used and 4 rooms remaining.'),
    ('Room Number', 'Read-only value generated automatically from the selected floor, for example P401.'),
    ('Area', 'Required room area in square meters.'),
    ('Room Type', 'Optional room classification, such as Studio or One Bedroom.'),
    ('Cancel', 'Closes the form without creating a room.'),
    ('Save New Room', 'Validates the fields, checks that the selected floor has available capacity, creates the room, and refreshes the building overview.'),
])

doc.add_paragraph('1.3 Create and Assign Assets', style='Heading 3')
add_body(doc, 'This feature manages the asset lifecycle from creating an asset record and storing it in the selected building inventory to assigning it to one room or multiple rooms. The system must keep each asset identifiable, prevent invalid duplicate assignments, and maintain the asset location and condition history.')

doc.add_paragraph('1.3.1 Asset Management Screen', style='Heading 4')
add_body(doc, 'This screen is the entry point for asset management. It presents assets grouped by type or by room and provides search, status, asset-class, and asset-type filters.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Page header', 'Shows the Asset Management title, current building context, total asset summary, and authorized actions.'),
    ('Create Asset', 'Opens the Create Asset form for creating a new asset type and one or more asset records.'),
    ('Restock Quantity', 'Opens the Restock form for adding more units to an existing asset model.'),
    ('Search field', 'Searches by asset name, room, asset code, model, or serial information.'),
    ('Condition filter', 'Filters assets by condition, such as Good, Damaged, Under Repair, or Inactive.'),
    ('Asset class filter', 'Filters by Fixed Asset, Equipment or Supplies, and Consumable.'),
    ('Asset type filter', 'Filters by a specific asset type, such as Air Conditioner, Refrigerator, Bed, or Smart Lock.'),
    ('View mode', 'Switches between assets grouped by type and assets grouped by room.'),
    ('Asset type card', 'Summarizes the number of assets, stock count, assigned-room count, condition badges, and quick assignment action.'),
    ('Room asset list', 'Shows assets currently assigned to each room and exposes detail, recovery, and condition actions.'),
])

doc.add_paragraph('1.3.2 Create Asset Screen', style='Heading 4')
add_body(doc, 'This form creates a new asset model and stores the requested quantity in the selected building inventory. The asset code is generated by the system and must not be manually edited.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Building context', 'Read-only display of the building selected on the top bar. The user does not select another building in this form.'),
    ('Asset type name', 'Required. Selects an existing type or enters a new type name.'),
    ('Asset class', 'Required. Classifies the asset as Fixed Asset, Equipment or Supplies, or Consumable.'),
    ('Asset code', 'Read-only preview of the generated identifier based on the asset type and building context.'),
    ('Brand', 'Required manufacturer or brand name.'),
    ('Model', 'Required model or product reference.'),
    ('Main specification', 'Required capacity or main specification, such as Inverter 1.5HP.'),
    ('Initial quantity', 'Required number of units to create in the inventory.'),
    ('Display name', 'Read-only generated name composed from the asset type, brand, specification, and model.'),
    ('Purchase price', 'Required original purchase price in VND.'),
    ('Warranty period', 'Required warranty duration in months.'),
    ('Create and Save to Inventory', 'Validates the form, generates unique asset codes for every unit, and saves the units to inventory.'),
])

doc.add_paragraph('1.3.3 Restock Quantity Screen', style='Heading 4')
add_body(doc, 'This screen adds additional units to an existing asset model without creating a duplicate model. The user first filters by asset class, then selects an asset type and a specific model.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Asset class', 'Required first filter. Limits the available asset types to the selected class.'),
    ('Asset type', 'Required. Lists asset types that already have inventory records in the selected class.'),
    ('Asset model', 'Required. Selects the existing model that will receive additional units.'),
    ('Additional quantity', 'Required positive quantity to add to inventory.'),
    ('Asset condition', 'Required condition for the new units, such as New 100%, Good, or Used.'),
    ('New purchase price', 'Required unit purchase price for the restock batch.'),
    ('Restock note', 'Optional note describing the procurement batch.'),
    ('Confirm', 'Creates unique asset records for the added quantity and stores them as available inventory.'),
])

doc.add_paragraph('1.3.4 Assign Asset to Room Screen', style='Heading 4')
add_body(doc, 'This screen assigns one available asset to a selected room in the current building. The user chooses the asset model, selects a room, records the handover note, and confirms the assignment.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Building context', 'Read-only selected building used to limit available inventory and destination rooms.'),
    ('Available asset', 'Required. Lists assets currently in stock and available for assignment.'),
    ('Destination room', 'Required. Lists rooms in the selected building. A quick filter can show only rooms that do not already have the selected fixed asset type.'),
    ('Assignment note', 'Optional note describing the handover condition or included accessories.'),
    ('Confirm assignment', 'Validates availability and duplicate rules, then changes the asset location from inventory to the selected room.'),
])

doc.add_paragraph('1.3.5 Quick Bulk Assignment Screen', style='Heading 4')
add_body(doc, 'This screen supports assigning the same asset type to multiple rooms in one operation. The system displays eligible rooms as selectable items and excludes rooms that already contain the same fixed asset type.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Asset type', 'Read-only type selected from the asset type card that initiated the quick assignment.'),
    ('Eligible room list', 'Displays rooms in the current building that can receive the selected asset type.'),
    ('Room selection', 'Allows the user to select one or more destination rooms.'),
    ('Assignment note', 'Optional note applied to the generated assignment records.'),
    ('Confirm bulk assignment', 'Assigns available units to all selected rooms or returns a clear error when stock is insufficient.'),
])

doc.add_paragraph('1.3.6 Asset Detail and Location Actions', style='Heading 4')
add_body(doc, 'The asset detail drawer shows the full record and assignment history for one asset. From this view, an authorized user can update condition, transfer an assigned asset to another room, or recover it to inventory.')
add_body(doc, 'Field Description')
add_table(doc, [
    ('Field Name', 'Description'),
    ('Asset identity', 'Shows asset code, display name, type, brand, model, specifications, purchase price, warranty, and current condition.'),
    ('Current location', 'Shows the current building and room, or the inventory location when the asset is unassigned.'),
    ('Assignment history', 'Shows assignment, transfer, and recovery events with room, date, user, and notes.'),
    ('Update condition', 'Changes the current condition and records the update in the audit history.'),
    ('Transfer room', 'Moves an assigned asset to another eligible room and records a transfer event.'),
    ('Recover to inventory', 'Removes the asset from its room, returns it to inventory, and records a recovery event.'),
])

doc.add_paragraph('1.3.7 Asset Processing Rules', style='Heading 4')
for item in [
    'Every created unit receives a unique asset code. The code is generated by the system and cannot be edited manually.',
    'A fixed asset type cannot be assigned twice to the same room. Equipment types configured to allow multiples may be assigned more than once.',
    'Only assets in available inventory can be assigned to a room.',
    'A bulk assignment must not partially create duplicate assignments; each result must be reported clearly when stock is insufficient.',
    'Every assignment, transfer, recovery, condition update, and restock operation must be recorded in asset history and audit logs.',
    'Staff users can perform only the asset actions allowed by their role; creating asset models and restocking are restricted from Staff.'
]:
    add_bullet(doc, item)

doc.add_paragraph('1.4 Building Status Management', style='Heading 3')
doc.add_paragraph('1.4.1 Deactivation Confirmation', style='Heading 4')
add_body(doc, 'A building can be temporarily deactivated from the building overview. If the building still contains rooms with active rental contracts, the system blocks the operation and shows the list of affected rooms. If no room is currently rented, the system asks for confirmation before changing the building status to Inactive.')

doc.add_paragraph('2. Roles and Permissions', style='Heading 2')
add_table(doc, [
    ('Role', 'Create Building flow permission'),
    ('Administrator', 'Can create, update, activate, and deactivate buildings and manage rooms.'),
    ('Manager', 'Can create and manage buildings and rooms within the operational scope.'),
    ('Staff', 'Can view building and room information and perform permitted operational actions, but cannot create a building.'),
    ('Resident', 'Cannot access building administration screens.'),
])

doc.add_paragraph('3. Background Job Inventory', style='Heading 2')
add_body(doc, 'No background job is required for the Create Building flow. Building and room generation is completed synchronously as part of the form submission.')

doc.add_paragraph('4. External API Inventory', style='Heading 2')
add_body(doc, 'No external API is required for the current Create Building prototype. Building, room, and validation data are handled by the application data store.')

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(OUTPUT)
