from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REFERENCE = Path(r"C:\Users\tienhop\Downloads\Report-3_Software Requirement Specification (4).docx")
OUTPUT = Path(r"D:\ki9\prototype\artifacts\Create_Contract.docx")
SCREEN_IMAGES = {
    '3.1.1 Contract Management Screen': (Path(r"C:\Users\tienhop\AppData\Local\Temp\codex-clipboard-ad9866eb-2456-43ad-aace-be6bb4a5ca72.png"), 'Figure 1. Contract Management Screen'),
    '3.1.2 Create Contract Screen': (Path(r"C:\Users\tienhop\AppData\Local\Temp\codex-clipboard-53a7408b-cdf6-4d84-899d-47194b0180d6.png"), 'Figure 2. Create Contract Screen'),
}


def clear_body(document):
    body = document._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_body(document, text):
    paragraph = document.add_paragraph(text, style='Normal')
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_screen_image(document, key):
    image_path, caption = SCREEN_IMAGES[key]
    if not image_path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
    caption_paragraph = document.add_paragraph(caption, style='Normal')
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_paragraph.runs:
        run.italic = True


def add_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.rows[0].cells[0].text = 'Field Name'
    table.rows[0].cells[1].text = 'Description'
    for field, description in rows:
        cells = table.add_row().cells
        cells[0].text = field
        cells[1].text = description

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), '9000')
    tbl_w.set(qn('w:type'), 'dxa')
    layout = tbl_pr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tbl_pr.append(layout)
    layout.set(qn('w:type'), 'fixed')

    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = borders.find(qn('w:' + edge))
        if border is None:
            border = OxmlElement('w:' + edge)
            borders.append(border)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')

    for row in table.rows:
        for cell, width in zip(row.cells, ('2500', '6500')):
            cell.width = Inches(int(width) / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), width)
            tc_w.set(qn('w:type'), 'dxa')
            for paragraph in cell.paragraphs:
                paragraph.style = document.styles['Normal']

    for run in table.rows[0].cells[0].paragraphs[0].runs + table.rows[0].cells[1].paragraphs[0].runs:
        run.bold = True
    return table


doc = Document(str(REFERENCE))
clear_body(doc)

sect_pr = doc._body._element.sectPr
pg_mar = sect_pr.find(qn('w:pgMar'))
if pg_mar is not None:
    for key, value in {'top': '720', 'bottom': '720', 'left': '720', 'right': '720', 'header': '360', 'footer': '360'}.items():
        pg_mar.set(qn('w:' + key), value)

doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(10)
for name in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = None

doc.add_paragraph('III. Functional Requirements', style='Heading 1')
add_body(doc, 'This feature manages the rental contract lifecycle from creating a contract for an eligible room and resident to activating, monitoring, settling, and closing the contract. The system must validate the room and resident relationship, preserve contract history, and prevent invalid overlapping contracts.')
doc.add_paragraph('3. Contract Management', style='Heading 2')

doc.add_paragraph('3.1 Contract Creation and Registration', style='Heading 3')
doc.add_paragraph('3.1.1 Contract Management Screen', style='Heading 4')
add_screen_image(doc, '3.1.1 Contract Management Screen')
add_body(doc, 'This screen is the contract-management workspace for the selected building. The header shows the page title, breadcrumb, building selector, and notification entry point. The action area provides Contract Template and Create New Contract. The contract list card contains status filters for All, Pending Approval, Active, and Rejected contracts, a keyword search field, and a table showing contract code, room and building, tenant, identity number, rental period, rent and deposit, status, and the View Contract action.')
doc.add_paragraph('Field Description', style='Normal')
add_table(doc, [
    ('Building selector', 'Sets the current building context and limits the displayed contracts to that building.'),
    ('Contract template', 'Opens the contract-template management flow for users with permission to edit the rental template.'),
    ('Create Contract', 'Opens the form for creating a new rental contract.'),
    ('Status chips', 'Filters the list by All, Pending approval, Active, or Rejected status and displays the count for each status.'),
    ('Search field', 'Searches by room, tenant name, or contract code.'),
    ('Contract table', 'Shows contract code, room and building, tenant, identity number, rental period, rent/deposit, status, and detail action.'),
    ('View contract', 'Opens the full contract document and the actions allowed by the contract status and current role.'),
])

doc.add_paragraph('3.1.2 Create Contract Screen', style='Heading 4')
add_screen_image(doc, '3.1.2 Create Contract Screen')
add_body(doc, 'This modal creates a rental contract and prepares the completed contract document. The header contains two tabs: the information form and the full contract preview. The form is divided into three visible sections: the lessor and building context, the representative tenant information, and the lease term and financial terms. Required fields are marked with an asterisk. The system generates the contract code, calculates the end date from the selected start date and duration, and lets the user continue with the Fill Full Contract action after validation.')
doc.add_paragraph('Field Description', style='Normal')
add_table(doc, [
    ('Form / Preview tabs', 'Switches between entering contract data and viewing the completed contract document before submission.'),
    ('Building and room', 'Shows the operating building, manager, hotline, and a required selector for an available room.'),
    ('Tenant identity', 'Captures required tenant name, identity number, phone, Zalo number, email, and permanent address; date of birth and identity-document details are also recorded.'),
    ('Contract code', 'Read-only unique code generated by the system for the new contract.'),
    ('Lease duration', 'Selects the rental duration, such as 6, 12, or 24 months.'),
    ('Start and end dates', 'The user selects the start date; the system calculates the read-only end date from the duration.'),
    ('Monthly rent', 'Required rental amount used for the contract and later billing.'),
    ('Security deposit', 'Required deposit amount recorded for later settlement.'),
    ('Monthly due day', 'Records the monthly payment deadline, restricted to the configured day range.'),
    ('Fill full contract', 'Validates the form and opens the completed contract document in the Preview tab.'),
])

doc.add_paragraph('3.2 Contract Processing Rules', style='Heading 3')
add_body(doc, 'The system applies the following business rules when creating and managing contracts:')
for rule in [
    'All required fields must be completed before the contract can be saved or submitted.',
    'A room cannot have more than one overlapping active contract for the same period.',
    'A resident must be valid and eligible for the selected building and contract workflow.',
    'The contract code is unique, generated by the system, and cannot be manually duplicated.',
    'The end date must be later than the start date, and financial amounts must be valid non-negative values.',
    'Creating or activating a contract updates the room status and refreshes the building room overview.',
    'Contract status transitions must follow the defined lifecycle; completed or terminated contracts cannot be edited as active contracts.',
    'Every creation, approval, update, renewal, termination, cancellation, and settlement action is recorded in the contract history and audit log.',
]:
    add_body(doc, '- ' + rule)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(OUTPUT)
