from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REFERENCE = Path(r"C:\Users\tienhop\Downloads\Report-3_Software Requirement Specification (4).docx")
OUTPUT = Path(r"D:\ki9\prototype\artifacts\Create_Account.docx")
SCREEN_IMAGES = {
    '4.1.1 Account Management Screen': (Path(r"C:\Users\tienhop\AppData\Local\Temp\codex-clipboard-7253c606-1a46-43fa-8677-4b34469a3153.png"), 'Figure 1. Account Management Screen'),
    '4.1.2 Create Account Screen': (Path(r"C:\Users\tienhop\AppData\Local\Temp\codex-clipboard-66f0b472-4733-4928-8782-a6b5331ecf59.png"), 'Figure 2. Create Account Screen'),
    '4.1.3 Account Creation Success Popup': (Path(r"C:\Users\tienhop\AppData\Local\Temp\codex-clipboard-5dbfd187-5872-4bf8-b11d-669417665769.png"), 'Figure 3. Account Creation Success Popup'),
}


def clear_body(document):
    body = document._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_body(document, text):
    paragraph = document.add_paragraph(text, style='Normal')
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_screen_image(document, key):
    image_path, caption = SCREEN_IMAGES[key]
    if not image_path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
    caption_paragraph = document.add_paragraph(caption, style='Normal')
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_paragraph.runs:
        run.italic = True


def add_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.rows[0].cells[0].text = 'Field Name'
    table.rows[0].cells[1].text = 'Description'
    for field, description in rows:
        cells = table.add_row().cells
        cells[0].text = field
        cells[1].text = description

    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement('w:tblW')
    tbl_w.set(qn('w:w'), '9000')
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_pr.append(tbl_w)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(layout)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        borders.append(border)
    tbl_pr.append(borders)

    for row in table.rows:
        for cell, width in zip(row.cells, ('2500', '6500')):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = OxmlElement('w:tcW')
            tc_w.set(qn('w:w'), width)
            tc_w.set(qn('w:type'), 'dxa')
            tc_pr.append(tc_w)
            for paragraph in cell.paragraphs:
                paragraph.style = document.styles['Normal']
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    return table


doc = Document(str(REFERENCE))
clear_body(doc)

sect_pr = doc._body._element.sectPr
pg_mar = sect_pr.find(qn('w:pgMar'))
if pg_mar is not None:
    for key, value in {'top': '720', 'bottom': '720', 'left': '720', 'right': '720', 'header': '360', 'footer': '360'}.items():
        pg_mar.set(qn('w:' + key), value)

doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(10)
for name in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = None

doc.add_paragraph('III. Functional Requirements', style='Heading 1')
add_body(doc, 'This feature manages user accounts and access permissions from creating an account record to assigning a role, limiting its building scope, activating or locking access, and maintaining account audit history. The system must keep login identities unique and apply role-based access control to every account action.')
doc.add_paragraph('4. Account Management', style='Heading 2')

doc.add_paragraph('4.1 Account Creation and Administration', style='Heading 3')
doc.add_paragraph('4.1.1 Account Management Screen', style='Heading 4')
add_screen_image(doc, '4.1.1 Account Management Screen')
add_body(doc, 'This screen is the administration workspace for user accounts and permissions. The header shows the page title, breadcrumb, and notification entry point. The action area shows the total account count and the Create Account action. The toolbar supports keyword search and filters by approval state or role. When a resident account is created from a contract and is waiting for access approval, the screen shows an alert banner, an Inactive status badge, and a Duyệt Active action in the corresponding table row.')
doc.add_paragraph('Field Description', style='Normal')
add_table(doc, [
    ('Building selector', 'Sets the current operational context and supports account scope management for the selected building.'),
    ('Account count', 'Displays the total number of user accounts in the system.'),
    ('Create Account', 'Opens the form for adding a new user account and assigning its role and access scope.'),
    ('Search field', 'Searches by full name, email address, or phone number.'),
    ('Account filters', 'Filters accounts by all, pending approval, Admin, Manager, Staff, or Resident role.'),
    ('Pending-account alert', 'Shows the number of resident accounts waiting for access approval and provides a shortcut to the pending list.'),
    ('Account table', 'Shows full name and title, login email, phone number, role, building scope, account status, and available actions.'),
    ('Account actions', 'Allows an authorized user to approve a pending resident account, edit account data, issue a temporary password, or lock and unlock access.'),
])

doc.add_paragraph('4.1.2 Create Account Screen', style='Heading 4')
add_screen_image(doc, '4.1.2 Create Account Screen')
add_body(doc, 'This modal creates a new operational user account. The form contains required fields for full name, login email, phone number, and Zalo number; a role selector limited to Manager and Staff; a required responsible-building selector; and an access-status selector. The user submits the form with Create Account. The password is not entered manually; it is generated by the system after successful creation. There is no job-title field in this flow, and required fields are marked with an asterisk.')
doc.add_paragraph('Field Description', style='Normal')
add_table(doc, [
    ('Account name', 'Required full name of the person who will use the account.'),
    ('Login email', 'Required unique email used as the account login identifier.'),
    ('Phone number', 'Required contact phone number for the account holder.'),
    ('Zalo number', 'Required Zalo contact number used to send account-related notifications.'),
    ('Role', 'Required access role. The Create Account form allows only Manager or Staff.'),
    ('Responsible building', 'Required. Links the new Manager or Staff account to one active operational building.'),
    ('Access status', 'Sets the initial status as Active, Inactive pending approval, or Locked.'),
    ('Create Account', 'Validates the form, generates a temporary password, creates the account, records an audit event, and opens the success popup.'),
])

doc.add_paragraph('4.1.3 Account Creation Success Popup', style='Heading 4')
add_screen_image(doc, '4.1.3 Account Creation Success Popup')
add_body(doc, 'After a new account is created successfully, the system opens a confirmation popup. The popup shows the login username, generated temporary password, assigned working building, and the Zalo number that will receive the notification. It displays the prepared welcome message containing the username, temporary password, working building, and first-login password-change instruction. The administrator can copy the password, send the message through the entered Zalo number, and then close the popup after the sent status is recorded.')
doc.add_paragraph('Field Description', style='Normal')
add_table(doc, [
    ('Account username', 'Shows the email address used to sign in.'),
    ('Temporary password', 'Shows the system-generated password for the first login and allows it to be copied.'),
    ('Working building', 'Shows the operational building assigned during account creation.'),
    ('Zalo recipient', 'Shows the Zalo number entered for the new account and used as the notification destination.'),
    ('Zalo message', 'Shows the prepared welcome message containing the username, temporary password, working building, and first-login password-change instruction.'),
    ('Send Zalo notification', 'Sends the prepared message to the entered Zalo number, updates the notification status, and records an audit event.'),
    ('Security notice', 'Requires the account holder to change the temporary password at the first login.'),
    ('Close confirmation', 'Remains unavailable until the Zalo notification is sent successfully; only then can the administrator close the popup.'),
])

doc.add_paragraph('4.2 Account Processing Rules', style='Heading 3')
add_body(doc, 'The system applies the following business rules when creating and managing accounts:')
for rule in [
    'Full name, login email, phone number, and role are required before an account can be saved.',
    'The Zalo number is required for account-related notifications and is stored with the account record.',
    'The login email must be unique, ignoring letter case; duplicate login identities are rejected.',
    'The Create Account form allows only Manager and Staff roles; Admin and Resident accounts are not created from this flow.',
    'The responsible building is required and must refer to an active operational building.',
    'The system generates a temporary password after successful creation; the administrator does not enter the password manually.',
    'The first-login security policy requires the account holder to change the temporary password.',
    'The account-creation notification must be sent to the Zalo number entered in the form and must include the username, temporary password, working building, and password-change instruction.',
    'The success popup cannot be closed before the Zalo notification is sent successfully.',
    'Only authorized administrators can create accounts, change roles, change building scope, lock accounts, or issue temporary passwords.',
    'A Resident account created from a pending contract remains Inactive until an authorized administrator approves access.',
    'A Staff user cannot approve or reject rental contracts and cannot access administrator-only account controls.',
    'Locked or inactive accounts cannot sign in until an authorized user changes the account status to Active.',
    'Every account creation, update, activation, lock, unlock, and password reset is recorded in the audit log.',
]:
    add_body(doc, '- ' + rule)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(OUTPUT)
